"""Excel / CSV / JSON exports with proper types, formats, frozen headers and filters."""
import csv
import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="0F4C5C")
HEAD_FONT = Font(bold=True, color="FFFFFF")
FLAG_FILL = PatternFill("solid", fgColor="FDE2E2")
USD = '"$"#,##0.00'


def _style(ws, widths, money_cols=(), date_cols=(), pct_cols=()):
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column in money_cols:
                c.number_format = USD
            elif c.column in date_cols:
                c.number_format = "yyyy-mm-dd"
            elif c.column in pct_cols:
                c.number_format = "0.0%"


def write_invoices(records, path_xlsx, path_csv, path_json):
    """records: list of (filename, extracted_rec, problems)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append(["Invoice No", "Vendor", "Customer", "Invoice Date", "Due Date", "Items", "Subtotal",
               "Tax Rate", "Tax", "Total", "Pages", "Check"])
    li = wb.create_sheet("Line Items")
    li.append(["Invoice No", "Line", "Description", "Qty", "Unit Price", "Amount"])
    fl = wb.create_sheet("Flags")
    fl.append(["Invoice No", "Problem", "Printed Total", "Expected Total", "Difference"])
    flat = []
    for name, r, problems in records:
        ws.append([r.get("invoice_no"), r.get("vendor"), r.get("customer"),
                   date.fromisoformat(r["invoice_date"]) if r.get("invoice_date") else None,
                   date.fromisoformat(r["due_date"]) if r.get("due_date") else None,
                   len(r["items"]), r.get("subtotal_c", 0) / 100, r.get("tax_rate_bp", 0) / 10000,
                   r.get("tax_c", 0) / 100, r.get("total_c", 0) / 100, r["pages"],
                   "OK" if not problems else ", ".join(problems)])
        if problems:
            for c in ws[ws.max_row]:
                c.fill = FLAG_FILL
            expected = (r.get("subtotal_c", 0) + r.get("tax_c", 0)) / 100
            for p in problems:
                fl.append([r.get("invoice_no"), p, r.get("total_c", 0) / 100, expected,
                           (r.get("total_c", 0) / 100) - expected])
        for n, it in enumerate(r["items"], 1):
            li.append([r.get("invoice_no"), n, it["desc"], it["qty"], it["unit_c"] / 100, it["amount_c"] / 100])
        flat.append({"file": name, "invoice_no": r.get("invoice_no"), "vendor": r.get("vendor"),
                     "customer": r.get("customer"), "invoice_date": r.get("invoice_date"),
                     "due_date": r.get("due_date"), "subtotal": r.get("subtotal_c", 0) / 100,
                     "tax": r.get("tax_c", 0) / 100, "total": r.get("total_c", 0) / 100,
                     "check": problems, "items": [{"description": i["desc"], "qty": i["qty"],
                                                   "unit_price": i["unit_c"] / 100, "amount": i["amount_c"] / 100}
                                                  for i in r["items"]]})
    _style(ws, [16, 26, 28, 13, 13, 7, 13, 9, 11, 13, 7, 22], money_cols=(7, 9, 10), date_cols=(4, 5), pct_cols=(8,))
    _style(li, [16, 6, 36, 7, 12, 13], money_cols=(5, 6))
    _style(fl, [16, 22, 14, 15, 12], money_cols=(3, 4, 5))
    wb.save(path_xlsx)
    with open(path_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["invoice_no", "vendor", "customer", "invoice_date", "due_date", "items", "subtotal", "tax",
                    "total", "check"])
        for d in flat:
            w.writerow([d["invoice_no"], d["vendor"], d["customer"], d["invoice_date"], d["due_date"],
                        len(d["items"]), d["subtotal"], d["tax"], d["total"], ";".join(d["check"]) or "OK"])
    with open(path_json, "w", encoding="utf-8") as f:
        json.dump(flat, f, ensure_ascii=False, indent=1)
    return {"invoices": ws.max_row - 1, "line_items": li.max_row - 1, "flags": fl.max_row - 1}


def write_statements(records, path_xlsx):
    """records: list of (name, extracted_rec, flagged_rows, summary_problems)."""
    wb = Workbook()
    sm = wb.active
    sm.title = "Summary"
    sm.append(["Statement", "Account", "Rows", "Opening", "Debits", "Credits", "Closing", "Reconciled"])
    for name, r, flagged, problems in records:
        s = r["summary"]
        sm.append([name.replace(".pdf", ""), r.get("account_no"), len(r["rows"]), s.get("opening_c", 0) / 100,
                   s.get("total_debits_c", 0) / 100, s.get("total_credits_c", 0) / 100,
                   s.get("closing_c", 0) / 100,
                   "Yes" if not flagged and not problems else f"No: {len(flagged)} row(s) flagged"])
        ws = wb.create_sheet(name.replace(".pdf", ""))
        ws.append(["Date", "Description", "Reference", "Debit", "Credit", "Balance"])
        for i, row in enumerate(r["rows"]):
            ws.append([date.fromisoformat(row["date"]), row["desc"], row["ref"],
                       row["debit_c"] / 100 or None, row["credit_c"] / 100 or None, row["balance_c"] / 100])
            if i in flagged:
                for c in ws[ws.max_row]:
                    c.fill = FLAG_FILL
        _style(ws, [12, 38, 14, 12, 12, 14], money_cols=(4, 5, 6), date_cols=(1,))
    _style(sm, [14, 18, 7, 13, 13, 13, 13, 26], money_cols=(4, 5, 6, 7))
    wb.save(path_xlsx)
    return {"sheets": len(wb.sheetnames), "rows": sum(len(r["rows"]) for _, r, _, _ in records)}
